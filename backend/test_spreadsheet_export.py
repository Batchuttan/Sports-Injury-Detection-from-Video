import io
import requests
import openpyxl

BASE_URL = "http://127.0.0.1:8000"

def test_exports():
    print("Testing Authentication...")
    login_res = requests.post(f"{BASE_URL}/token", data={"username": "demo@sportsai.com", "password": "password123"})
    assert login_res.status_code == 200, "Login failed"
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Single Video Analysis Export
    print("\n--- 1. Testing Single Analysis XLSX Export ---")
    analyses_res = requests.get(f"{BASE_URL}/athlete/analyses", headers=headers).json()
    assert len(analyses_res) > 0, "No analyses found"
    analysis_id = analyses_res[0]["id"]
    activity = analyses_res[0]["activity"]
    print(f"Testing Analysis #{analysis_id} ({activity})...")

    xlsx_res = requests.get(f"{BASE_URL}/athlete/analyses/{analysis_id}/export/csv?format=xlsx", headers=headers)
    assert xlsx_res.status_code == 200, f"XLSX export failed: {xlsx_res.status_code}"
    assert "spreadsheetml.sheet" in xlsx_res.headers.get("Content-Type", "")
    
    # Load workbook from memory
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    ws = wb.active
    print(f"Sheet title: {ws.title}")
    print(f"Max rows: {ws.max_row}, Max cols: {ws.max_column}")

    # Inspect Column E dimensions
    col_e_width = ws.column_dimensions["E"].width
    print(f"Column E width: {col_e_width}")
    assert col_e_width is not None and col_e_width >= 25.0, f"Expected Col E width >= 25, got {col_e_width}"
    print(f"-> SUCCESS: Column E width is {col_e_width} (well above 13!)")

    # Inspect Cell 6 in Column E (Row 6, Col 5)
    cell_6_val = ws.cell(row=6, column=5).value
    print(f"Row 6 Column E (Reference) value: '{cell_6_val}'")
    assert cell_6_val == "Action Detection", f"Expected 'Action Detection', got '{cell_6_val}'"
    print(f"-> SUCCESS: 'Action Detection' is present in Row 6 Col E with width {col_e_width}")

    # Check other column widths
    print(f"Column A width: {ws.column_dimensions['A'].width}")
    print(f"Column B width: {ws.column_dimensions['B'].width}")
    print(f"Column C width: {ws.column_dimensions['C'].width}")
    print(f"Column D width: {ws.column_dimensions['D'].width}")
    print(f"Column F width: {ws.column_dimensions['F'].width}")
    assert ws.column_dimensions["B"].width >= 40.0
    assert ws.column_dimensions["F"].width >= 80.0

    # Check wrap text on row 6 column 6
    cell_f6 = ws.cell(row=6, column=6)
    assert cell_f6.alignment.wrap_text == True
    print(f"-> SUCCESS: Column F has wrap_text={cell_f6.alignment.wrap_text}")

    # 2. Dual Comparison Export
    print("\n--- 2. Testing Dual Comparison XLSX Export ---")
    comp_res = requests.get(f"{BASE_URL}/athlete/comparisons", headers=headers).json()
    if comp_res:
        comp_id = comp_res[0]["id"]
        print(f"Testing Dual Comparison #{comp_id}...")
        dual_xlsx = requests.get(f"{BASE_URL}/athlete/comparisons/{comp_id}/export/csv?format=xlsx", headers=headers)
        if dual_xlsx.status_code != 200:
            print(f"Error {dual_xlsx.status_code}: {dual_xlsx.text}")
        assert dual_xlsx.status_code == 200, f"Dual export failed: {dual_xlsx.status_code} {dual_xlsx.text}"
        wb_dual = openpyxl.load_workbook(io.BytesIO(dual_xlsx.content))
        ws_dual = wb_dual.active
        print(f"Dual Sheet: {ws_dual.title}, Max rows: {ws_dual.max_row}")
        for col_let in ["A", "B", "C", "D", "E", "F"]:
            w = ws_dual.column_dimensions[col_let].width
            print(f"Dual Col {col_let} width: {w}")
            assert w >= 20.0
        print("-> SUCCESS: Dual comparison XLSX has generous column widths across all columns!")

    # 3. Test Raw CSV with BOM
    print("\n--- 3. Testing Raw CSV fallback ---")
    csv_res = requests.get(f"{BASE_URL}/athlete/analyses/{analysis_id}/export/csv?format=csv", headers=headers)
    assert csv_res.status_code == 200
    assert csv_res.content.startswith(b"\xef\xbb\xbf"), "Missing UTF-8 BOM in CSV"
    print("-> SUCCESS: Raw CSV includes UTF-8 BOM for Microsoft Excel encoding support.")

    print("\n==========================================")
    print("ALL EXPORT WIDTH VERIFICATIONS PASSED!")
    print("==========================================")

if __name__ == "__main__":
    test_exports()
